# agentic-document-automator/create_sample_documents.py
"""
สร้างไฟล์ input สำหรับ Document Agent:
- templates/report_template.docx : template Word ที่มี placeholder {{KEY}} และ {{TABLE_PLACEHOLDER}}
- documents/source_document.pdf  : เอกสาร 2 หน้า สำหรับใส่ลายน้ำและดึงข้อความ
- documents/dummy_doc1.pdf, dummy_doc2.pdf : เอกสารสำหรับทดสอบการรวม PDF
"""
import os

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DOCS_DIR = os.path.join(BASE_DIR, "documents")
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")
styles = getSampleStyleSheet()


def create_template():
    doc = docx.Document()
    doc.add_heading("{{REPORT_TITLE}}", level=0)

    meta = doc.add_paragraph()
    meta.add_run("Date: ").bold = True
    meta.add_run("{{REPORT_DATE}}")
    meta.add_run("    Author: ").bold = True
    meta.add_run("{{AUTHOR_NAME}}")
    org = doc.add_paragraph()
    org.add_run("Organization: ").bold = True
    # placeholder ที่ถูกแบ่งเป็น 2 run (สถานการณ์ที่เกิดบ่อยเมื่อแก้ template ใน Word)
    org.add_run("{{COMPANY_")
    org.add_run("NAME}}")

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph("{{INTRODUCTION_TEXT}}")

    doc.add_heading("2. Monthly Overview: {{MONTH_PERIOD}}", level=1)
    doc.add_paragraph("Here is a summary of key performance indicators:")
    doc.add_paragraph("{{TABLE_PLACEHOLDER}}")

    doc.add_heading("3. Conclusion", level=1)
    doc.add_paragraph("This report concludes the performance analysis for {{MONTH_PERIOD}}.")

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "{{COMPANY_NAME}} - generated automatically by Document Agent"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    path = os.path.join(TEMPLATES_DIR, "report_template.docx")
    doc.save(path)
    print(f"Created {os.path.relpath(path, BASE_DIR)}")


def build_pdf(filename, pages):
    path = os.path.join(DOCS_DIR, filename)
    story = []
    for i, (heading, paragraphs) in enumerate(pages):
        story.append(Paragraph(heading, styles["Title"] if i == 0 else styles["Heading1"]))
        story.append(Spacer(1, 12))
        for text in paragraphs:
            story.append(Paragraph(text, styles["BodyText"]))
            story.append(Spacer(1, 8))
        if i < len(pages) - 1:
            story.append(PageBreak())
    SimpleDocTemplate(path, pagesize=A4, title=filename).build(story)
    print(f"Created {os.path.relpath(path, BASE_DIR)} ({len(pages)} page(s))")


if __name__ == "__main__":
    os.makedirs(DOCS_DIR, exist_ok=True)
    os.makedirs(TEMPLATES_DIR, exist_ok=True)
    create_template()
    build_pdf("source_document.pdf", [
        ("Internal Project Proposal", [
            "Project: Agentic Document Automation. This document describes an internal tool that "
            "generates reports, protects PDF files with watermarks and merges documentation.",
            "Scope: Word report generation from templates, PDF watermarking, PDF merging and text extraction.",
            "Owner: Aussadawut Ruenkaew. Status: Approved for internal use only.",
        ]),
        ("Budget and Timeline", [
            "Phase 1 (Week 1): Template design and configuration format.",
            "Phase 2 (Week 2): PDF tasks (watermark, merge, extract).",
            "Phase 3 (Week 3): Audit report and testing. Estimated budget: 0 THB (open-source libraries).",
        ]),
    ])
    build_pdf("dummy_doc1.pdf", [
        ("Project Documentation - Part 1", ["Overview of the Document Agent architecture: config_parser, "
                                            "document_agent, pdf_tasks, word_tasks and utils."]),
    ])
    build_pdf("dummy_doc2.pdf", [
        ("Project Documentation - Part 2", ["User guide: edit configs/document_automation_config.json "
                                            "and run python main.py."]),
    ])
