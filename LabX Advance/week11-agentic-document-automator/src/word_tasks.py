# agentic-document-automator/src/word_tasks.py
import os

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

from utils import setup_logging

logger = setup_logging(__name__)

TABLE_PLACEHOLDER = "{{TABLE_PLACEHOLDER}}"


class WordTasks:
    """Handles advanced Word document operations for the Document Agent."""

    def __init__(self):
        logger.info("WordTasks initialized.")

    @staticmethod
    def _replace_in_paragraph(paragraph, replacements):
        """
        แทนที่ placeholder ในย่อหน้า แม้ placeholder จะถูก Word แบ่งอยู่หลาย run
        (เช่น '{{REPORT_' กับ 'TITLE}}') โดยรวมข้อความทุก run แล้วเขียนกลับลง run แรก
        คงรูปแบบตัวอักษร (bold/size/color) ของ run แรกไว้
        """
        full_text = "".join(run.text for run in paragraph.runs)
        if "{{" not in full_text:
            return 0
        new_text = full_text
        count = 0
        for key, value in replacements.items():
            token = "{{" + key + "}}"
            if token in new_text:
                count += new_text.count(token)
                new_text = new_text.replace(token, str(value))
        if count:
            paragraph.runs[0].text = new_text
            for run in paragraph.runs[1:]:
                run.text = ""
        return count

    def _insert_table_at(self, document, anchor_paragraph, table_data):
        """สร้างตารางแล้วย้ายไปวางตรงตำแหน่งย่อหน้า placeholder (แทนที่จะต่อท้ายเอกสาร)"""
        cols = len(table_data[0])
        table = document.add_table(rows=0, cols=cols)
        try:
            table.style = "Table Grid"
        except Exception:
            logger.warning("Style 'Table Grid' not available, using default table style.")
        for r_idx, row_data in enumerate(table_data):
            cells = table.add_row().cells
            for c_idx, value in enumerate(row_data):
                cells[c_idx].text = str(value)
                if r_idx == 0:
                    for run in cells[c_idx].paragraphs[0].runs:
                        run.bold = True
        anchor_paragraph._p.addnext(table._tbl)
        anchor_paragraph._p.getparent().remove(anchor_paragraph._p)
        return table

    def generate_report_from_template(self, template_path, output_path, data_placeholders, table_data=None):
        """
        สร้างรายงาน Word จาก template:
        - แทนที่ {{KEY}} ด้วยค่าจาก data_placeholders (ทั้งในย่อหน้า ตาราง header และ footer)
        - แทรกตาราง table_data (แถวแรกเป็นหัวตาราง) ที่ตำแหน่ง {{TABLE_PLACEHOLDER}}
        """
        if not os.path.exists(template_path):
            logger.error(f"Word template not found: {template_path}")
            return False
        try:
            document = docx.Document(template_path)

            anchor = None
            for p in document.paragraphs:
                if TABLE_PLACEHOLDER in p.text:
                    anchor = p
                    continue
                self._replace_in_paragraph(p, data_placeholders)

            containers = list(document.tables)
            for section in document.sections:
                containers += [section.header, section.footer]
            for container in containers:
                rows = container.rows if hasattr(container, "rows") else None
                if rows is not None:
                    for row in rows:
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                self._replace_in_paragraph(p, data_placeholders)
                else:
                    for p in container.paragraphs:
                        self._replace_in_paragraph(p, data_placeholders)

            if table_data:
                if anchor is not None:
                    self._insert_table_at(document, anchor, table_data)
                    logger.info(f"Inserted {len(table_data)}x{len(table_data[0])} table at placeholder.")
                else:
                    logger.warning(f"'{TABLE_PLACEHOLDER}' not found in template. Appending table to end.")
                    tmp = document.add_paragraph()
                    self._insert_table_at(document, tmp, table_data)

            leftovers = [p.text for p in document.paragraphs if "{{" in p.text]
            if leftovers:
                logger.warning(f"Unfilled placeholders remain: {leftovers}")

            os.makedirs(os.path.dirname(output_path), exist_ok=True)
            document.save(output_path)
            logger.info(f"Report generated from template and saved to '{output_path}'.")
            return True
        except PermissionError:
            logger.error(f"Cannot write '{output_path}'. Is the file open in Word?")
            return False
        except Exception as e:
            logger.error(f"Error generating report from template: {e}")
            return False

    def extract_text_from_word(self, docx_path):
        """Extracts all text (paragraphs + tables in document order) from a Word document."""
        if not os.path.exists(docx_path):
            logger.error(f"Word document not found: {docx_path}")
            return None
        try:
            document = docx.Document(docx_path)
            lines = []
            body = document.element.body
            for child in body.iterchildren():
                tag = child.tag.split("}")[-1]
                if tag == "p":
                    lines.append(Paragraph(child, document).text)
                elif tag == "tbl":
                    table = Table(child, document)
                    for row in table.rows:
                        lines.append("\t".join(cell.text for cell in row.cells))
            logger.info(f"Text extracted from Word document: {docx_path}")
            return "\n".join(lines)
        except Exception as e:
            logger.error(f"Error extracting text from Word document '{docx_path}': {e}")
            return None
