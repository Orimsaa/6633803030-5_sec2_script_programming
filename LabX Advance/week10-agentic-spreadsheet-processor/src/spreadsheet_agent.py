# agentic-spreadsheet-processor/src/spreadsheet_agent.py
import datetime
import json

import openpyxl
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation

from utils import setup_logging

logger = setup_logging(__name__)

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


class SpreadsheetAgent:
    """
    Agent ที่อ่าน "แผนงาน" (tasks) จาก configuration แล้วลงมือทำกับไฟล์ Excel ทีละขั้นตอน
    (perception = input + tasks, action = แก้ไขเวิร์กบุ๊ก) พร้อมบันทึกทุกขั้นตอนลงชีต "Audit Log"
    """

    def __init__(self, config):
        self.config = config
        self.workbook = None
        self.audit_log = []
        self.start_time = datetime.datetime.now()
        # ตาราง dispatch: task type -> method
        self.handlers = {
            "copy_data": self._process_copy_data_task,
            "calculate_column": self._process_calculate_column_task,
            "conditional_format": self._process_conditional_format_task,
            "chart": self._process_chart_task,
            "summary_sheet": self._process_summary_sheet_task,
            "data_validation": self._process_data_validation_task,
        }
        logger.info("Spreadsheet Agent initialized.")

    # ------------------------------------------------------------------ helpers
    def _log_audit(self, status, message, task_type="N/A", details=None):
        self.audit_log.append({
            "timestamp": datetime.datetime.now().isoformat(timespec="seconds"),
            "status": status,
            "task_type": task_type,
            "message": message,
            "details": details if details else {},
        })
        if status in ("ERROR", "CRITICAL"):
            logger.error(f"AUDIT LOG - {status}: {message} | Task: {task_type}")
        elif status == "WARNING":
            logger.warning(f"AUDIT LOG - {status}: {message} | Task: {task_type}")
        else:
            logger.info(f"AUDIT LOG - {status}: {message} | Task: {task_type}")

    def _load_workbook(self):
        input_file = self.config["input_file"]
        try:
            self.workbook = openpyxl.load_workbook(input_file)
            self._log_audit("SUCCESS", f"Workbook '{input_file}' loaded.", "load_workbook",
                            {"sheets": self.workbook.sheetnames})
            return True
        except FileNotFoundError:
            self._log_audit("ERROR", f"Input file not found: {input_file}", "load_workbook")
        except Exception as e:
            self._log_audit("ERROR", f"Failed to load workbook: {e}", "load_workbook")
        return False

    def _get_sheet(self, sheet_name, create=True):
        if sheet_name in self.workbook.sheetnames:
            return self.workbook[sheet_name]
        if not create:
            raise KeyError(f"Sheet '{sheet_name}' not found")
        logger.info(f"Creating new sheet '{sheet_name}'.")
        return self.workbook.create_sheet(sheet_name)

    @staticmethod
    def _last_data_row(sheet, column="A"):
        """แถวสุดท้ายที่มีข้อมูลในคอลัมน์ (ไม่นับแถวว่างท้ายชีต)"""
        for r in range(sheet.max_row, 0, -1):
            if sheet[f"{column}{r}"].value not in (None, ""):
                return r
        return 1

    @staticmethod
    def _style_header_row(sheet, columns):
        for c in range(1, columns + 1):
            cell = sheet.cell(row=1, column=c)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")

    @staticmethod
    def _autofit(sheet, min_width=10, max_width=60):
        for col_cells in sheet.columns:
            letter = get_column_letter(col_cells[0].column)
            longest = max((len(str(c.value)) for c in col_cells if c.value is not None), default=0)
            sheet.column_dimensions[letter].width = max(min_width, min(longest + 2, max_width))

    # ------------------------------------------------------------------ tasks
    def _process_copy_data_task(self, task):
        src_name, dest_name = task.get("source_sheet"), task.get("dest_sheet")
        if not src_name or not dest_name:
            raise ValueError("copy_data requires 'source_sheet' and 'dest_sheet'")
        src = self._get_sheet(src_name, create=False)
        dest = self._get_sheet(dest_name)
        rows = 0
        for row in src.iter_rows():
            if all(c.value in (None, "") for c in row):
                continue
            rows += 1
            for cell in row:
                dest.cell(row=rows, column=cell.column, value=cell.value)
        self._style_header_row(dest, src.max_column)
        self._autofit(dest)
        return f"Copied {rows - 1} data rows from '{src_name}' to '{dest_name}'"

    def _process_calculate_column_task(self, task):
        sheet_name, target_col, formula = task.get("sheet"), task.get("target_column"), task.get("formula")
        if not all([sheet_name, target_col, formula]):
            raise ValueError("calculate_column requires 'sheet', 'target_column' and 'formula'")
        sheet = self._get_sheet(sheet_name)
        start_row = task.get("start_row", 2)
        last_row = self._last_data_row(sheet)
        header = sheet[f"{target_col}1"]
        header.value = task.get("header", "Calculated Column")
        header.font, header.fill, header.alignment = HEADER_FONT, HEADER_FILL, Alignment(horizontal="center")
        for r in range(start_row, last_row + 1):
            cell = sheet[f"{target_col}{r}"]
            cell.value = formula.replace("{row}", str(r))
            if task.get("number_format"):
                cell.number_format = task["number_format"]
        sheet.column_dimensions[target_col].width = max(14, len(header.value) + 4)
        return f"Column '{target_col}' ({header.value}) filled with formula for rows {start_row}-{last_row}"

    def _process_conditional_format_task(self, task):
        sheet = self._get_sheet(task["sheet"])
        rng, rule_type = task.get("range"), task.get("rule_type")
        if rule_type == "color_scale":
            rule = ColorScaleRule(start_type="min", start_color=task.get("min_color", "F8696B"),
                                  mid_type="percentile", mid_value=50, mid_color=task.get("mid_color", "FFEB84"),
                                  end_type="max", end_color=task.get("max_color", "63BE7B"))
        elif rule_type == "data_bar":
            rule = DataBarRule(start_type="min", end_type="max", color=task.get("bar_color", "638EC6"),
                               showValue=True)
        elif rule_type == "expression":
            fill_color = task.get("fill_color", "FFC7CE")
            dxf = DifferentialStyle(font=Font(color=task.get("font_color", "9C0006"), bold=True),
                                    fill=PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid"))
            rule = Rule(type="expression", formula=[task["formula"].lstrip("=")], dxf=dxf)
        else:
            raise ValueError(f"Unsupported conditional format rule_type: {rule_type}")
        sheet.conditional_formatting.add(rng, rule)
        return f"Applied '{rule_type}' formatting to {task['sheet']}!{rng}"

    def _process_chart_task(self, task):
        sheet = self._get_sheet(task["sheet"])
        chart_type = task.get("chart_type", "bar")
        charts = {"bar": BarChart, "line": LineChart, "pie": PieChart}
        if chart_type not in charts:
            raise ValueError(f"Unsupported chart type: {chart_type}")
        chart = charts[chart_type]()
        chart.title = task.get("title", "Chart")
        if chart_type != "pie":
            chart.x_axis.title = task.get("x_axis_title")
            chart.y_axis.title = task.get("y_axis_title")
            chart.legend.position = "b"
        chart.width, chart.height = task.get("width", 18), task.get("height", 9)
        min_col, min_row, max_col, max_row = range_boundaries(task["data_range"])
        data = Reference(sheet, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
        chart.add_data(data, titles_from_data=True)
        if task.get("category_range"):
            c1, r1, c2, r2 = range_boundaries(task["category_range"])
            chart.set_categories(Reference(sheet, min_col=c1, min_row=r1, max_col=c2, max_row=r2))
        sheet.add_chart(chart, task.get("top_left_cell", "H2"))
        return f"{chart_type.title()} chart '{task.get('title', 'Chart')}' placed at {task.get('top_left_cell', 'H2')}"

    def _process_summary_sheet_task(self, task):
        """สร้างชีตสรุปผลรวม/ค่าเฉลี่ย/สูงสุด ด้วยสูตร Excel ที่อ้างอิงข้อมูลจากชีตต้นทาง"""
        sheet = self._get_sheet(task.get("sheet", "Summary"))
        source = self._get_sheet(task["source_sheet"], create=False)
        last = self._last_data_row(source)
        src = f"'{source.title}'"
        sheet.append(["Metric", "Value"])
        for item in task["metrics"]:
            formula = item["formula"].replace("{src}", src).replace("{last}", str(last))
            sheet.append([item["label"], formula])
            if item.get("number_format"):
                sheet.cell(row=sheet.max_row, column=2).number_format = item["number_format"]
        self._style_header_row(sheet, 2)
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=2):
            for cell in row:
                cell.border = BORDER
        sheet.column_dimensions["A"].width = 30
        sheet.column_dimensions["B"].width = 20
        return f"Summary sheet '{sheet.title}' created with {len(task['metrics'])} metrics from rows 2-{last}"

    def _process_data_validation_task(self, task):
        """เพิ่ม dropdown (Data Validation แบบ list) ให้คอลัมน์ พร้อมค่าเริ่มต้น"""
        sheet = self._get_sheet(task["sheet"])
        col = task["target_column"]
        last = self._last_data_row(sheet)
        header = sheet[f"{col}1"]
        header.value = task.get("header", "Validated")
        header.font, header.fill, header.alignment = HEADER_FONT, HEADER_FILL, Alignment(horizontal="center")
        options = task["options"]
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Invalid value",
                            error=f"Please choose one of: {', '.join(options)}")
        sheet.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")
        for r in range(2, last + 1):
            sheet[f"{col}{r}"].value = task.get("default", options[0])
        sheet.column_dimensions[col].width = max(16, len(header.value) + 4)
        return f"Dropdown {options} added to {sheet.title}!{col}2:{col}{last}"

    # ------------------------------------------------------------------ report
    def _add_audit_log_sheet(self):
        audit = self.workbook.create_sheet("Audit Log", 0)
        audit.append(["Timestamp", "Status", "Task Type", "Message", "Details"])
        for c in range(1, 6):
            cell = audit.cell(row=1, column=c)
            cell.font = HEADER_FONT
            cell.fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
        colors = {"SUCCESS": "C6EFCE", "ERROR": "FFC7CE", "CRITICAL": "FFC7CE", "WARNING": "FFEB9C"}
        for entry in self.audit_log:
            audit.append([entry["timestamp"], entry["status"], entry["task_type"], entry["message"],
                          json.dumps(entry["details"], ensure_ascii=False)])
            color = colors.get(entry["status"])
            if color:
                audit.cell(row=audit.max_row, column=2).fill = PatternFill(start_color=color, end_color=color,
                                                                           fill_type="solid")
        self._autofit(audit)
        audit.freeze_panes = "A2"
        logger.info("Audit log sheet created.")

    def run(self):
        if not self._load_workbook():
            self._log_audit("CRITICAL", "Agent could not start", "run")
            return False
        ok = True
        for i, task in enumerate(self.config["tasks"], 1):
            task_type = task.get("type")
            logger.info(f"Executing Task {i}: Type='{task_type}'")
            handler = self.handlers.get(task_type)
            if handler is None:
                self._log_audit("WARNING", f"Unsupported task type: {task_type}", "task_execution", task)
                continue
            try:
                message = handler(task)
                self._log_audit("SUCCESS", message, task_type)
            except Exception as e:
                ok = False
                self._log_audit("ERROR", f"Task {i} failed: {e}", task_type, task)

        duration = (datetime.datetime.now() - self.start_time).total_seconds()
        self._log_audit("INFO", f"Agent finished {len(self.config['tasks'])} tasks in {duration:.2f} seconds.",
                        "agent_run_summary")
        try:
            self._add_audit_log_sheet()
            output_file = self.config["output_file"]
            self.workbook.save(output_file)
            logger.info(f"Final workbook saved to: {output_file} (sheets: {self.workbook.sheetnames})")
        except PermissionError:
            logger.critical(f"Cannot write '{self.config['output_file']}'. Close it in Excel and run again.")
            return False
        except Exception as e:
            logger.critical(f"Failed to save output workbook: {e}", exc_info=True)
            return False
        return ok
